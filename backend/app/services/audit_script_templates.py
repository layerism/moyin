import json
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile


PYTHON_TEMPLATE = """import json
import sys
from pathlib import Path

import fitz
import openpyxl
from docx import Document
from PIL import Image
from pptx import Presentation


ALLOWED_EXTENSIONS = {".docx", ".xlsx", ".pdf", ".pptx", ".jpeg", ".jpg", ".png"}


def validate_payload(payload: dict) -> list[dict]:
    if not isinstance(payload, dict) or payload.get("schemaVersion") != "1.0":
        raise ValueError("不支持的输入协议版本")
    files = payload.get("files")
    if not isinstance(files, list) or not files:
        raise ValueError("files 必须是非空数组")
    for file in files:
        if not isinstance(file, dict):
            raise ValueError("文件信息格式错误")
        extension = str(file.get("extension", "")).lower()
        path = Path(str(file.get("path", "")))
        if extension not in ALLOWED_EXTENSIONS or path.suffix.lower() != extension:
            raise ValueError("文件扩展名不受支持或与路径不一致")
        if not path.is_absolute() or not path.is_file():
            raise ValueError("文件路径无效")
    return files


def parse_file(file: dict) -> dict:
    path = Path(file["path"])
    extension = file["extension"].lower()
    if extension == ".docx":
        document = Document(path)
        return {"paragraphCount": len(document.paragraphs)}
    if extension == ".xlsx":
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return {"sheetNames": workbook.sheetnames}
        finally:
            workbook.close()
    if extension == ".pdf":
        with fitz.open(path) as document:
            return {"pageCount": document.page_count}
    if extension == ".pptx":
        presentation = Presentation(path)
        return {"slideCount": len(presentation.slides)}
    with Image.open(path) as image:
        image.verify()
        return {"format": image.format}


def run(payload: dict) -> dict:
    files = validate_payload(payload)
    issues = []
    for file in files:
        try:
            parse_file(file)
        except Exception:
            issues.append({
                "fileId": str(file.get("id", "")),
                "code": "FILE_PARSE_ERROR",
                "message": "文件解析失败",
            })
    return {
        "schemaVersion": "1.0",
        "passed": not issues,
        "reason": "" if not issues else "部分材料解析失败",
        "details": {"checkedFileCount": len(files), "issues": issues},
    }


def main() -> None:
    try:
        payload = json.loads(sys.stdin.read())
        print(json.dumps(run(payload), ensure_ascii=False))
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        raise


if __name__ == "__main__":
    main()
"""

JAVASCRIPT_TEMPLATE = r"""const { readFileSync } = require("node:fs");
const { readFile } = require("node:fs/promises");
const mammoth = require("mammoth");
const readExcelFile = require("read-excel-file/node");
const { PDFParse } = require("pdf-parse");
const JSZip = require("jszip");
const { XMLParser } = require("fast-xml-parser");
const sharp = require("sharp");

const ALLOWED_EXTENSIONS = new Set([".docx", ".xlsx", ".pdf", ".pptx", ".jpeg", ".jpg", ".png"]);

function validatePayload(payload) {
  if (!payload || payload.schemaVersion !== "1.0") throw new Error("不支持的输入协议版本");
  if (!Array.isArray(payload.files) || payload.files.length === 0) throw new Error("files 必须是非空数组");
  for (const file of payload.files) {
    const extension = String(file.extension || "").toLowerCase();
    const path = String(file.path || "");
    if (!ALLOWED_EXTENSIONS.has(extension) || !path.toLowerCase().endsWith(extension)) {
      throw new Error("文件扩展名不受支持或与路径不一致");
    }
  }
  return payload.files;
}

async function parseFile(file) {
  const extension = file.extension.toLowerCase();
  if (extension === ".docx") {
    const result = await mammoth.extractRawText({ path: file.path });
    return { characterCount: result.value.length };
  }
  if (extension === ".xlsx") {
    const sheets = await readExcelFile(file.path);
    return { sheetNames: sheets.map((sheet) => sheet.sheet) };
  }
  if (extension === ".pdf") {
    const parser = new PDFParse({ data: await readFile(file.path) });
    try {
      const result = await parser.getText();
      return { pageCount: result.total };
    } finally {
      await parser.destroy();
    }
  }
  if (extension === ".pptx") {
    const archive = await JSZip.loadAsync(await readFile(file.path));
    const parser = new XMLParser();
    const slides = Object.keys(archive.files).filter((name) => /^ppt\/slides\/slide\d+\.xml$/.test(name));
    for (const name of slides) parser.parse(await archive.file(name).async("string"));
    return { slideCount: slides.length };
  }
  return sharp(file.path).metadata();
}

async function run(payload) {
  const files = validatePayload(payload);
  const issues = [];
  for (const file of files) {
    try {
      await parseFile(file);
    } catch {
      issues.push({
        fileId: String(file.id || ""),
        code: "FILE_PARSE_ERROR",
        message: "文件解析失败",
      });
    }
  }
  return {
    schemaVersion: "1.0",
    passed: issues.length === 0,
    reason: issues.length === 0 ? "" : "部分材料解析失败",
    details: { checkedFileCount: files.length, issues },
  };
}

async function main() {
  try {
    const payload = JSON.parse(readFileSync(0, "utf8"));
    process.stdout.write(JSON.stringify(await run(payload)));
  } catch (error) {
    process.stderr.write(`${error instanceof Error ? error.message : String(error)}\n`);
    process.exitCode = 1;
  }
}

void main();
"""

INPUT_EXAMPLE = json.dumps(
    {
        "schemaVersion": "1.0",
        "executionId": "00000000-0000-0000-0000-000000000000",
        "files": [
            {
                "id": "file-1",
                "name": "申请表.docx",
                "extension": ".docx",
                "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "path": "/tmp/audit/files/file-1.docx",
                "size": 1024,
                "sha256": "example-sha256",
            }
        ],
        "context": {
            "workflowId": "workflow-id",
            "flowInstanceId": "instance-id",
            "nodeId": "node-id",
            "submitterId": "user-id",
        },
    },
    ensure_ascii=False,
    indent=2,
)

OUTPUT_EXAMPLE = json.dumps(
    {
        "schemaVersion": "1.0",
        "passed": True,
        "reason": "",
        "details": {"checkedFileCount": 1, "issues": []},
    },
    ensure_ascii=False,
    indent=2,
)

README_BY_LANGUAGE = {
    "python": """# Python 审核脚本模板\n\n只修改 `handler.py`，最终上传该文件。运行环境预装 python-docx、openpyxl、PyMuPDF、python-pptx、Pillow。脚本从 stdin 读取一个 JSON 对象，stdout 只能输出最终 JSON，诊断日志写入 stderr。不得在脚本内安装依赖、访问 OSS 或依赖项目环境变量。支持 `.docx .xlsx .pdf .pptx .jpeg .jpg .png`。\n""",
    "javascript": """# JavaScript 审核脚本模板\n\n只修改 `handler.js`，最终上传该文件。运行环境预装 read-excel-file、mammoth、pdf-parse、jszip、fast-xml-parser、sharp。脚本从 stdin 读取一个 JSON 对象，stdout 只能输出最终 JSON，诊断日志写入 stderr。不得在脚本内安装依赖、访问 OSS 或依赖项目环境变量。支持 `.docx .xlsx .pdf .pptx .jpeg .jpg .png`。\n""",
}


def get_template_archive(language: str) -> tuple[bytes, str]:
    if language == "python":
        entry_name = "handler.py"
        entry_source = PYTHON_TEMPLATE
        filename = "audit-script-python-template.zip"
    elif language == "javascript":
        entry_name = "handler.js"
        entry_source = JAVASCRIPT_TEMPLATE
        filename = "audit-script-javascript-template.zip"
    else:
        raise ValueError("仅支持 Python 或 JavaScript 模板")

    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr(entry_name, entry_source)
        archive.writestr("input.example.json", INPUT_EXAMPLE)
        archive.writestr("output.example.json", OUTPUT_EXAMPLE)
        archive.writestr("README.md", README_BY_LANGUAGE[language])
    return output.getvalue(), filename
